from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .models import DraftDocument


def export_docx(draft: DraftDocument, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    document.add_heading(draft.title, level=0)

    for block in sorted(draft.blocks, key=lambda item: item.order):
        content = block.content
        if block.type == "heading":
            document.add_heading(str(content.get("text", "")), level=clamp_heading(content.get("level", 2)))
        elif block.type == "paragraph":
            document.add_paragraph(str(content.get("text", "")))
        elif block.type == "table":
            add_caption(document, content.get("caption", "표"))
            add_table(document, content.get("headers", []), content.get("rows", []))
        elif block.type == "chart":
            add_caption(document, content.get("title", "차트"))
            add_table(
                document,
                ["항목", "값"],
                [[item.get("label", ""), item.get("value", "")] for item in content.get("values", [])],
            )
            if content.get("note"):
                document.add_paragraph(str(content["note"]))
        elif block.type == "diagram":
            add_caption(document, content.get("title", "다이어그램"))
            add_code_block(document, content.get("code", ""))
        elif block.type == "image":
            document.add_paragraph(str(content.get("alt", "이미지")))

    document.save(output_path)
    return output_path


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)


def clamp_heading(value: Any) -> int:
    try:
        return min(4, max(1, int(value)))
    except (TypeError, ValueError):
        return 2


def add_caption(document: Document, text: object) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(text or ""))
    run.bold = True


def add_table(document: Document, headers: list[object], rows: list[list[object]]) -> None:
    column_count = max(len(headers), *(len(row) for row in rows), 1)
    table = document.add_table(rows=1 if headers else 0, cols=column_count)
    table.style = "Table Grid"
    if headers:
        for index in range(column_count):
            table.rows[0].cells[index].text = stringify(headers[index] if index < len(headers) else "")
    for row in rows:
        cells = table.add_row().cells
        for index in range(column_count):
            cells[index].text = stringify(row[index] if index < len(row) else "")
    document.add_paragraph()


def add_code_block(document: Document, code: object) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(str(code or ""))
    run.font.name = "Courier New"
    run.font.size = Pt(9)


def stringify(value: object) -> str:
    return "" if value is None else str(value)
