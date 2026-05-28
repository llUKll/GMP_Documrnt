from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any


def parse_file(path: Path, filename: str) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        return parse_xlsx(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".zip":
        return parse_zip(path, filename)
    if suffix == ".hwp":
        return parse_hwp_fallback(path, filename)
    if suffix == ".hwpx":
        return parse_hwpx(path, filename)
    if suffix in {".txt", ".md", ".csv", ".kt", ".java", ".xml", ".gradle", ".kts", ".json", ".yml", ".yaml"}:
        return parse_text(path, filename)
    raise ValueError(f"지원하지 않는 파일 형식입니다: {suffix}")


def parse_text(path: Path, filename: str) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8", errors="replace")[:20000]
    return {"kind": "text", "filename": filename, "text": text}


def parse_hwp_fallback(path: Path, filename: str) -> dict[str, Any]:
    return {
        "kind": "hwp",
        "filename": filename,
        "text": (
            f"[HWP: {filename}]\n"
            "HWP 바이너리 한글 문서입니다. 서버 모드에서도 본문 직접 추출은 제한되어 "
            "파일명/크기 메타데이터를 근거 목록에 포함하고, 본문 기반 문구는 [확인 필요]로 표시합니다.\n"
            f"파일 크기: {path.stat().st_size} bytes"
        ),
    }


def parse_hwpx(path: Path, filename: str) -> dict[str, Any]:
    import re
    import zipfile

    parts: list[str] = [f"[HWPX: {filename}]"]
    with zipfile.ZipFile(path) as package:
        names = [name for name in package.namelist() if name.lower().endswith(".xml")]
        for name in [n for n in names if re.search(r"contents?|section|body|header|footer", n, re.I)][:30]:
            raw = package.read(name).decode("utf-8", errors="replace")
            text = re.sub(r"<[^>]+>", " ", raw)
            text = " ".join(text.split())
            if text:
                parts.append(f"[XML: {name}]\n{text[:5000]}")
    return {"kind": "hwpx", "filename": filename, "text": "\n\n".join(parts)}


def parse_zip(path: Path, filename: str) -> dict[str, Any]:
    import re
    import zipfile

    def ignored(name: str) -> bool:
        normalized = name.replace("\\", "/")
        return bool(re.search(r"(^|/)\.git/|(^|/)node_modules/|(^|/)build/|(^|/)\.gradle/|(^|/)__macosx/|\.DS_Store$", normalized, re.I))

    def priority(name: str) -> tuple[int, str]:
        lower = name.lower()
        if re.search(r"ib-sdf|sdf|상세설계|software.*design|설계|인허가|regulatory|submission", lower):
            return (0, name)
        if re.search(r"readme|manual|guide|사용자|매뉴얼|보고서|report|log|test|검증|시험", lower):
            return (1, name)
        if re.search(r"androidmanifest\.xml$|build\.gradle|settings\.gradle", lower):
            return (2, name)
        if re.search(r"res/values/.*\.xml$|strings\.xml$|res/layout/.*\.xml$|res/menu/.*\.xml$|navigation/.*\.xml$", lower):
            return (3, name)
        if re.search(r"activity|service|repository|database|export|log|safety|alarm|bluetooth|ble|socket|api", lower):
            return (4, name)
        return (9, name)

    text_ext = {".txt", ".md", ".csv", ".kt", ".java", ".xml", ".gradle", ".kts", ".json", ".yml", ".yaml"}
    doc_ext = {".docx", ".pdf", ".xlsx", ".hwpx", ".hwp", ".zip"}
    with zipfile.ZipFile(path) as package:
        names = [name for name in package.namelist() if not name.endswith("/") and not ignored(name)]
        docs = [name for name in names if Path(name).suffix.lower() in doc_ext | text_ext]
        media = [name for name in names if re.search(r"\.(png|jpe?g|webp|gif|bmp|svg|mp4|mov|webm|avi|mkv)$", name, re.I)]
        important = sorted(docs, key=priority)[:120]
        parts = [
            f"[ZIP: {filename}]",
            "ZIP 분석 요약: 내부 문서, 코드/설정, 화면/미디어 파일명을 분류했습니다.",
            f"총 파일 수: {len(names)}",
            f"내부 문서/텍스트 후보: {len(docs)}",
            f"이미지/영상 후보: {len(media)}",
            "",
            "[주요 내부 파일]",
            *[f"- {name}" for name in important[:80]],
        ]
        if media:
            parts.extend(["", "[화면/미디어 후보]", *[f"- {name}" for name in media[:50]]])
        extracted = 0
        for name in important:
            if extracted >= 60 or sum(len(part) for part in parts) > 40000:
                break
            suffix = Path(name).suffix.lower()
            try:
                if suffix in text_ext:
                    data = package.read(name).decode("utf-8", errors="replace")
                    text = "\n".join(line.rstrip() for line in data.splitlines()).strip()[:7000]
                elif suffix == ".hwp":
                    info = package.getinfo(name)
                    text = f"[HWP: {name}]\nHWP 바이너리 문서입니다. 본문 추출 제한. 크기: {info.file_size} bytes"
                else:
                    continue
                if text:
                    parts.append(f"\n[ZIP FILE: {name}]\n{text}")
                    extracted += 1
            except Exception as exc:  # pragma: no cover - defensive parsing
                parts.append(f"\n[ZIP FILE: {name}]\n읽기 실패: {exc}")
        return {"kind": "zip", "filename": filename, "text": "\n".join(parts), "file_count": len(names)}


def parse_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        rows = [
            [cell for cell in row]
            for row in worksheet.iter_rows(values_only=True)
            if any(cell is not None and str(cell).strip() != "" for cell in row)
        ]
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else f"컬럼 {idx + 1}" for idx, value in enumerate(rows[0])]
        data_rows = rows[1:51]
        columns = infer_columns(headers, data_rows)
        sheets.append(
            {
                "name": worksheet.title,
                "headers": headers,
                "rows": [[normalize_cell(value) for value in row] for row in data_rows],
                "row_count": max(0, len(rows) - 1),
                "columns": columns,
            }
        )
    return {"kind": "xlsx", "sheets": sheets}


def infer_columns(headers: list[str], rows: list[list[Any]]) -> list[dict[str, Any]]:
    columns: list[dict[str, Any]] = []
    for index, header in enumerate(headers):
        values = [row[index] for row in rows if index < len(row) and row[index] is not None]
        numeric_values = [float(value) for value in values if isinstance(value, (int, float))]
        date_like = sum(1 for value in values if hasattr(value, "isoformat"))
        categories = Counter(str(value) for value in values if not isinstance(value, (int, float)))
        if numeric_values:
            kind = "number"
        elif date_like:
            kind = "date"
        else:
            kind = "category"
        columns.append(
            {
                "name": header,
                "kind": kind,
                "non_empty": len(values),
                "sum": round(sum(numeric_values), 2) if numeric_values else None,
                "avg": round(sum(numeric_values) / len(numeric_values), 2) if numeric_values else None,
                "top_values": categories.most_common(5),
            }
        )
    return columns


def normalize_cell(value: Any) -> Any:
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def parse_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows[:30]:
            rows.append([cell.text.strip() for cell in row.cells])
        if rows:
            tables.append(rows)
    return {
        "kind": "docx",
        "paragraphs": paragraphs[:80],
        "tables": tables[:10],
        "image_count": len(document.inline_shapes),
    }


def parse_pdf(path: Path) -> dict[str, Any]:
    from PyPDF2 import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page_index, page in enumerate(reader.pages[:20], start=1):
        text = page.extract_text() or ""
        pages.append({"page": page_index, "text": text[:4000]})
    return {"kind": "pdf", "pages": pages, "page_count": len(reader.pages)}
