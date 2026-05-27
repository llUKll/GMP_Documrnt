from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from openpyxl import Workbook
from PyPDF2 import PdfWriter

from app.analyzer import build_draft
from app.docx_export import export_docx
from app.hwpx_export import export_hwpx
from app.parsers import parse_docx, parse_pdf, parse_xlsx


class MvpPipelineTest(unittest.TestCase):
    def test_xlsx_parse_infers_numeric_columns_and_generates_chart(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sales.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "매출"
            sheet.append(["월", "매출", "고객군"])
            sheet.append(["1월", 120, "A"])
            sheet.append(["2월", 180, "B"])
            workbook.save(path)

            extracted = parse_xlsx(path)
            extracted["filename"] = "sales.xlsx"
            draft = build_draft("p1", "매출 분석", [extracted])

            self.assertEqual(extracted["sheets"][0]["columns"][1]["kind"], "number")
            self.assertTrue(any(block.type == "table" for block in draft.blocks))
            self.assertTrue(any(block.type == "chart" for block in draft.blocks))
            self.assertTrue(any(block.type == "diagram" for block in draft.blocks))

    def test_docx_and_pdf_parse_fail_soft_inputs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "reference.docx"
            document = Document()
            document.add_heading("분석 배경", level=1)
            document.add_paragraph("팀 내부 검토용 문서입니다.")
            document.save(docx_path)

            pdf_path = Path(tmp) / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=200)
            with pdf_path.open("wb") as target:
                writer.write(target)

            self.assertIn("팀 내부", " ".join(parse_docx(docx_path)["paragraphs"]))
            self.assertEqual(parse_pdf(pdf_path)["page_count"], 1)

    def test_hwpx_export_creates_package(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            draft = build_draft("p1", "문서", [])
            output = export_hwpx(draft, Path(tmp) / "document.hwpx")

            self.assertTrue(output.exists())
            with zipfile.ZipFile(output) as package:
                names = set(package.namelist())
                self.assertIn("Contents/content.hpf", names)
                self.assertIn("Contents/section0.xml", names)
                section = package.read("Contents/section0.xml").decode("utf-8")
                self.assertIn("문서 분석 보고서 초안", section)

    def test_docx_export_creates_word_document(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            draft = build_draft("p1", "문서", [])
            output = export_docx(draft, Path(tmp) / "document.docx")

            self.assertTrue(output.exists())
            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("문서 분석 보고서", text)


if __name__ == "__main__":
    unittest.main()
